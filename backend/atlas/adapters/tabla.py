"""Adaptador `tabla`: .csv, .tsv, .xlsx. Hojas → secciones; columnas → elementos `column`.

El tipo de dato se infiere desde los valores (no solo del nombre). Se conservan solo
estadísticas agregadas por columna (n, nulos, distintos, ejemplo de formato): nunca filas.
Regla permanente M4: sin datos de pacientes en el grafo.
"""
from __future__ import annotations

import csv
from pathlib import Path

from ..infer.rules import infer_data_type_from_values
from ..ir import IRDocument, IRElement, IRSection

NAME = "tabla@1.0"
EXTENSIONS = (".csv", ".tsv", ".xlsx")
SOURCE_TYPE = "tabla"
MAX_ROWS = 5000


def parse(path: Path) -> IRDocument:
    doc = IRDocument(filename=path.name, title=path.stem, source_type=SOURCE_TYPE, adapter=NAME)
    if path.suffix.lower() == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        for si, ws in enumerate(wb.worksheets):
            rows = [[("" if v is None else str(v)) for v in r] for r in ws.iter_rows(values_only=True, max_row=MAX_ROWS + 1)]
            doc.sections.append(_sheet_section(ws.title, rows, {"sheet": ws.title, "index": si}))
    else:
        delim = "\t" if path.suffix.lower() == ".tsv" else None
        with open(path, newline="", encoding="utf-8-sig") as f:
            sample = f.read(4096); f.seek(0)
            if delim is None:
                try:
                    delim = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
                except csv.Error:
                    delim = ","
            rows = [r for _, r in zip(range(MAX_ROWS + 1), csv.reader(f, delimiter=delim))]
        doc.sections.append(_sheet_section(path.stem, rows, {"file": path.name}))
    return doc


def _sheet_section(name: str, rows: list[list[str]], locator: dict) -> IRSection:
    sec = IRSection(name=name, locator=locator)
    if not rows:
        return sec
    header, body = rows[0], rows[1:]
    sec.description = f"{len(body)} filas leídas (máximo {MAX_ROWS}); solo se conservan estadísticas por columna."
    for ci, col in enumerate(header):
        col = (col or "").strip()
        if not col:
            continue
        values = [r[ci] if ci < len(r) else "" for r in body]
        non_empty = [v for v in values if str(v).strip() != ""]
        inf = infer_data_type_from_values(non_empty, col)
        distinct = {str(v).strip() for v in non_empty}
        stats = {"rows": len(values), "non_null": len(non_empty), "null_pct": round(100 * (1 - len(non_empty) / len(values)), 1) if values else None,
                 "distinct": len(distinct), "format_example": _format_example(non_empty)}
        el = IRElement(name=col, raw=col, locator={**locator, "column": ci}, element_kind="column",
                       declared={}, stats=stats)
        el.declared["_data_type_from_values"] = {"value": inf.value, "confidence": inf.confidence, "rationale": inf.rationale}
        # dominio de valores propuesto: solo si es categórica o binaria con pocos valores (sin filas)
        if inf.value in ("categorica", "binario") and len(distinct) <= 30:
            el.declared["_value_domain"] = sorted(distinct)[:30]
        sec.elements.append(el)
    return sec


def _format_example(vals: list[str]) -> str | None:
    """Ejemplo de *formato* (no de valor): dígitos → 9, letras → a. Evita filtrar datos."""
    if not vals:
        return None
    v = str(vals[0])[:24]
    out = []
    for ch in v:
        out.append("9" if ch.isdigit() else "a" if ch.isalpha() else ch)
    return "".join(out)
